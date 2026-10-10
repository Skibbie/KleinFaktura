"""Import bestehender Rechnungen/Angebote aus .ods, .xlsx und .pdf – reines Python.

- .ods  : zipfile + xml.etree (content.xml), kein LibreOffice nötig
- .xlsx : openpyxl
- .pdf  : pypdf (Textextraktion, bestmöglich)

Verarbeitet werden ausschließlich die Dateien, die der Benutzer in der GUI auswählt.
Der Parser ist auf das bisherige Vorlagenlayout abgestimmt (Pos./Beschreibung/Anzahl/Einheit/
Einzelpreis/Gesamtpreis, „Summe:“), arbeitet aber tolerant – alle Felder werden im Prüfdialog bestätigt.
"""
from __future__ import annotations

import re
import shutil
import sqlite3
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from xml.etree import ElementTree as ET

from app import config
from app.db import repo
from app.db.database import transaction
from app.money import line_total_cents, parse_date, to_cents, to_decimal
from app.services import audit, documents, numbering

SUPPORTED = (".ods", ".xlsx", ".xlsm", ".pdf")

ODS_NS = {
    "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
    "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
    "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
}
RX_NUMBER = re.compile(r"\b((?:AN)?\d{4}-\d{4,6})\b")
RX_DATE = re.compile(r"\b(\d{1,2}\.\d{1,2}\.\d{2,4})\b")
RX_ZIP_CITY = re.compile(r"^(?:D-)?(\d{5})\s+(.+)$")
RX_MONEY = re.compile(r"-?\d{1,3}(?:\.\d{3})*,\d{2}|-?\d+(?:[.,]\d{1,2})?")
SALUTATION_START = ("sehr geehrte", "hallo", "guten tag", "liebe", "lieber", "moin")


@dataclass
class ParsedLine:
    title: str = ""
    description: str = ""
    quantity: str = "1"
    unit: str = ""
    unit_price_cents: int = 0
    total_cents: int = 0


@dataclass
class ParsedInvoice:
    path: Path
    sha256: str
    doc_type: str = "rechnung"
    number: str = ""
    doc_date: str | None = None
    service_date: str | None = None
    salutation: str = ""
    recipient_lines: list[str] = field(default_factory=list)
    customer: dict = field(default_factory=dict)
    lines: list[ParsedLine] = field(default_factory=list)
    total_cents: int | None = None
    warnings: list[str] = field(default_factory=list)

    @property
    def lines_total(self) -> int:
        return sum(l.total_cents for l in self.lines)


# ---------------------------------------------------------------- Rohdaten lesen
def _ods_cell_text(cell) -> str:
    paras = []
    for p in cell.iter():
        if p.tag in (f"{{{ODS_NS['text']}}}p", f"{{{ODS_NS['text']}}}h"):
            parts = []
            for node in p.iter():
                if node is not p and node.tag == f"{{{ODS_NS['text']}}}s":
                    parts.append(" " * int(node.get(f"{{{ODS_NS['text']}}}c", "1")))
                elif node is not p and node.tag == f"{{{ODS_NS['text']}}}tab":
                    parts.append("\t")
                elif node is not p and node.tag == f"{{{ODS_NS['text']}}}line-break":
                    parts.append("\n")
                if node.text and (node is p or node.tag != f"{{{ODS_NS['text']}}}p"):
                    parts.append(node.text)
                if node is not p and node.tail:
                    parts.append(node.tail)
            paras.append("".join(parts))
    return "\n".join(paras).strip()


def read_ods(path: Path) -> list[list[str]]:
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read("content.xml"))
    t = ODS_NS["table"]
    rows_out: list[list[str]] = []
    sheet = root.find(f".//{{{t}}}table")
    if sheet is None:
        return rows_out
    for row in sheet.iter(f"{{{t}}}table-row"):
        cells: list[str] = []
        for cell in row:
            if cell.tag not in (f"{{{t}}}table-cell", f"{{{t}}}covered-table-cell"):
                continue
            rep = int(cell.get(f"{{{t}}}number-columns-repeated", "1"))
            val = _ods_cell_text(cell)
            vtype = cell.get(f"{{{ODS_NS['office']}}}value-type")
            if vtype in ("float", "currency", "percentage") and cell.get(f"{{{ODS_NS['office']}}}value") and not val:
                val = cell.get(f"{{{ODS_NS['office']}}}value")
            if vtype == "date" and cell.get(f"{{{ODS_NS['office']}}}date-value"):
                val = val or cell.get(f"{{{ODS_NS['office']}}}date-value")
            cells.extend([val] * (min(rep, 50) if val else min(rep, 1)))
        rrep = int(row.get(f"{{{t}}}number-rows-repeated", "1"))
        if any(cells):
            rows_out.extend([cells] * min(rrep, 50))
        else:
            rows_out.append([])
        if len(rows_out) > 2000:
            break
    return rows_out


def read_xlsx(path: Path) -> list[list[str]]:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    out = []
    for r in ws.iter_rows(values_only=True):
        row = []
        for v in r:
            if v is None:
                row.append("")
            elif hasattr(v, "strftime"):
                row.append(v.strftime("%d.%m.%Y"))
            elif isinstance(v, float):
                row.append(f"{v:.2f}".replace(".", ",") if abs(v - round(v, 2)) < 1e-9 else str(v))
            else:
                row.append(str(v).strip())
        out.append(row)
        if len(out) > 2000:
            break
    wb.close()
    return out


def read_pdf(path: Path) -> list[list[str]]:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    rows = []
    for page in reader.pages[:5]:
        text = page.extract_text() or ""
        for line in text.splitlines():
            line = line.replace("\u0000", "").strip()
            if line:
                rows.append([line])
    return rows


def read_rows(path: Path) -> list[list[str]]:
    ext = path.suffix.lower()
    if ext == ".ods":
        return read_ods(path)
    if ext in (".xlsx", ".xlsm"):
        return read_xlsx(path)
    if ext == ".pdf":
        return read_pdf(path)
    raise ValueError(f"Nicht unterstütztes Format: {ext}")


# ---------------------------------------------------------------- Parser
def _money_cents(s: str) -> int | None:
    s = (s or "").replace("€", "").replace("EUR", "").strip()
    if not s:
        return None
    m = RX_MONEY.findall(s)
    if not m:
        return None
    try:
        return to_cents(m[-1])
    except ValueError:
        return None


def _nonempty(row) -> list[str]:
    return [c for c in row if c and c.strip()]


def _is_own_company(text: str, company: dict) -> bool:
    t = text.lower()
    for key in ("name", "owner", "street", "phone", "email", "web", "tax_number", "iban", "line1"):
        v = (company.get(key) or "").lower()
        if v and (v in t or t in v) and len(t) > 3:
            return True
    return "steuernr" in t or "iban" in t or "@" in t or "www." in t


def parse_rows(rows: list[list[str]], path: Path, company: dict) -> ParsedInvoice:
    import hashlib
    inv = ParsedInvoice(path=path, sha256=hashlib.sha256(path.read_bytes()).hexdigest())
    flat = [" ".join(_nonempty(r)) for r in rows]
    alltext = "\n".join(flat)

    # Belegart
    low = alltext.lower()
    if "angebot" in low and "rechnung" not in low:
        inv.doc_type = "angebot"

    # Nummer: neben "Rechn…"/"Angebot"/"Nr", sonst erste passende, sonst Dateiname
    m = re.search(r"(?:rech\w*|angebot\w*|nr\.?|nummer)\s*[:.]?\s*((?:AN)?\d{4}-\d{4,6})", alltext, re.I)
    if m:
        inv.number = m.group(1)
    else:
        m = RX_NUMBER.search(alltext) or RX_NUMBER.search(path.stem)
        inv.number = m.group(1) if m else ""
    if not inv.number:
        inv.warnings.append("Keine Belegnummer erkannt.")

    # Anrede + Datum
    sal_idx = None
    for i, r in enumerate(rows):
        cells = _nonempty(r)
        if cells and cells[0].lower().startswith(SALUTATION_START):
            sal_idx = i
            inv.salutation = cells[0] if "," in cells[0] or len(cells) > 1 else cells[0]
            mm = RX_DATE.search(" ".join(cells))
            if mm:
                inv.doc_date = parse_date(mm.group(1))
            break
    if not inv.doc_date:
        for line in flat:
            mm = RX_DATE.search(line)
            if mm and parse_date(mm.group(1)):
                inv.doc_date = parse_date(mm.group(1))
                break
    if not inv.doc_date:
        inv.warnings.append("Kein Datum erkannt.")
    if inv.salutation and RX_DATE.search(inv.salutation):
        inv.salutation = RX_DATE.sub("", inv.salutation).strip()
    if "leistungsdatum" in low and ("entspricht" in low or "gleich" in low):
        inv.service_date = inv.doc_date
    else:
        mm = re.search(r"leistungs(?:datum|zeitraum)\s*:?\s*(\d{1,2}\.\d{1,2}\.\d{2,4})", alltext, re.I)
        inv.service_date = parse_date(mm.group(1)) if mm else inv.doc_date

    # Empfänger: Block direkt vor der Anrede (bzw. vor der ersten PLZ-Zeile)
    zip_idx = None
    search_end = sal_idx if sal_idx is not None else min(len(rows), 40)
    for i in range(search_end - 1, -1, -1):
        cells = _nonempty(rows[i])
        if cells and RX_ZIP_CITY.match(cells[0]) and not _is_own_company(cells[0], company):
            zip_idx = i
            break
    if zip_idx is not None:
        block = []
        i = zip_idx
        while i >= 0 and len(block) < 5:
            cells = _nonempty(rows[i])
            if not cells:
                if block:
                    break
                i -= 1
                continue
            if _is_own_company(cells[0], company) and block:
                break
            block.insert(0, cells[0])
            i -= 1
        inv.recipient_lines = block
        inv.customer = recipient_to_customer(block, inv.salutation)
    else:
        inv.warnings.append("Empfängeradresse nicht erkannt.")

    # Positionstabelle
    hdr_idx, cols = None, {}
    for i, r in enumerate(rows):
        lc = [c.lower() for c in r]
        if any(c.startswith("pos") for c in lc) and any("beschreibung" in c or "bezeichnung" in c or "leistung" in c for c in lc):
            hdr_idx = i
            for ci, c in enumerate(lc):
                if c.startswith("pos"):
                    cols.setdefault("pos", ci)
                elif "beschreibung" in c or "bezeichnung" in c or c == "leistung":
                    cols.setdefault("desc", ci)
                elif "anzahl" in c or "menge" in c:
                    cols.setdefault("qty", ci)
                elif "einheit" in c:
                    cols.setdefault("unit", ci)
                elif "einzel" in c:
                    cols.setdefault("price", ci)
                elif "gesamt" in c or "betrag" in c:
                    cols.setdefault("total", ci)
            break

    total_idx = None
    for i, r in enumerate(rows):
        joined = " ".join(_nonempty(r)).lower()
        if re.search(r"\b(summe|gesamtbetrag|rechnungsbetrag|endbetrag)\b", joined) and (hdr_idx is None or i > hdr_idx):
            vals = [_money_cents(c) for c in _nonempty(r)]
            vals = [v for v in vals if v is not None]
            if vals:
                inv.total_cents = vals[-1]
                total_idx = i
                if "gesamtbetrag" in joined or "endbetrag" in joined:
                    break

    if hdr_idx is not None and len(cols) >= 3 and len(rows[hdr_idx]) > 1:
        end = total_idx if total_idx is not None else len(rows)
        pending_desc: list[str] = []
        for r in rows[hdr_idx + 1:end]:
            def col(name):
                ci = cols.get(name)
                return r[ci].strip() if ci is not None and ci < len(r) and r[ci] else ""
            desc = col("desc")
            total = _money_cents(col("total"))
            if desc:
                pending_desc.append(desc)
            if total is not None and (total != 0 or pending_desc):
                text = "\n".join(pending_desc).strip()
                parts = text.split("\n", 1)
                qty = col("qty") or "1"
                price = _money_cents(col("price"))
                try:
                    qty_d = to_decimal(qty)
                except ValueError:
                    qty_d = to_decimal(1)
                if price is None:
                    price = int(total / qty_d) if qty_d else total
                pl = ParsedLine(title=parts[0], description=parts[1] if len(parts) > 1 else "",
                                quantity=str(qty_d), unit=col("unit"), unit_price_cents=price,
                                total_cents=line_total_cents(qty_d, price))
                if pl.total_cents != total:
                    inv.warnings.append(f"Position „{pl.title[:30]}“: Menge × Preis ≠ Gesamtpreis – bitte prüfen.")
                    pl.total_cents = total
                inv.lines.append(pl)
                pending_desc = []
    elif path.suffix.lower() == ".pdf":
        inv.lines = _parse_pdf_lines(rows)

    if not inv.lines and inv.total_cents:
        inv.lines.append(ParsedLine(title="Leistung laut Originalbeleg", quantity="1", unit="pauschal",
                                    unit_price_cents=inv.total_cents, total_cents=inv.total_cents))
        inv.warnings.append("Positionen nicht erkannt – Sammelposition mit Gesamtbetrag angelegt.")
    if inv.total_cents is None:
        inv.total_cents = inv.lines_total
        inv.warnings.append("Gesamtsumme nicht erkannt – aus Positionen berechnet.")
    elif inv.lines and inv.lines_total != inv.total_cents:
        inv.warnings.append("Summe der Positionen weicht von der Gesamtsumme ab – bitte prüfen.")
    return inv


_RX_PDF_ITEM = re.compile(
    r"^(\d{1,3})\s+(.*?)\s*(-?\d+(?:,\d+)?)\s+([A-Za-zÄÖÜäöüß.²]+)\s+(-?[\d.]+,\d{2})\s*€?\s+(-?[\d.]+,\d{2})\s*€?$")


def _parse_pdf_lines(rows: list[list[str]]) -> list[ParsedLine]:
    out = []
    for r in rows:
        m = _RX_PDF_ITEM.match(r[0].strip())
        if m:
            qty = to_decimal(m.group(3))
            price = to_cents(m.group(5))
            out.append(ParsedLine(title=m.group(2).strip() or "Position", quantity=str(qty), unit=m.group(4),
                                  unit_price_cents=price, total_cents=to_cents(m.group(6))))
    return out


def recipient_to_customer(block: list[str], salutation: str = "") -> dict:
    c: dict = {"source": "import"}
    lines = list(block)
    if lines:
        m = RX_ZIP_CITY.match(lines[-1])
        if m:
            c["zip"], c["city"] = m.group(1), m.group(2).strip()
            lines.pop()
    if lines:
        c["street"] = lines.pop()
    name_lines = [l for l in lines if l]
    person = name_lines[-1] if name_lines else ""
    if len(name_lines) >= 2:
        c["company"] = name_lines[0]
    for s in ("Herrn ", "Herr ", "Frau "):
        if person.startswith(s):
            c["salutation"] = s.strip().rstrip("n") if s.startswith("Herrn") else s.strip()
            person = person[len(s):]
    parts = person.split()
    if parts:
        c["last_name"] = parts[-1]
        c["first_name"] = " ".join(parts[:-1])
    sal = (salutation or "").lower()
    if not c.get("salutation"):
        if "frau" in sal:
            c["salutation"] = "Frau"
        elif "herr" in sal:
            c["salutation"] = "Herr"
    if salutation:
        c["letter_salutation"] = salutation if salutation.endswith(",") else salutation + ","
    return c


def parse_file(path: Path | str, company: dict) -> ParsedInvoice:
    path = Path(path)
    rows = read_rows(path)
    return parse_rows(rows, path, company)


# ---------------------------------------------------------------- Kollisionsprüfung
@dataclass
class Collision:
    kind: str          # hash / number / number_format / customer
    message: str
    existing: dict | None = None


def check_collisions(con: sqlite3.Connection, inv: ParsedInvoice, number: str | None = None,
                     doc_type: str | None = None) -> list[Collision]:
    number = (number if number is not None else inv.number).strip()
    doc_type = doc_type or inv.doc_type
    out = []
    ex = con.execute("SELECT * FROM document WHERE import_sha256=?", (inv.sha256,)).fetchone()
    if ex:
        out.append(Collision("hash", f"Diese Datei wurde bereits importiert (Beleg {ex['number']}).", dict(ex)))
    if number:
        ex = con.execute("SELECT * FROM document WHERE number=?", (number,)).fetchone()
        if ex:
            out.append(Collision("number", f"Die Nummer {number} ist bereits vergeben "
                                           f"({documents.DOC_TYPE_LABEL[ex['doc_type']]} vom {ex['doc_date']}, "
                                           f"{ex['total_gross_cents'] / 100:.2f} €).".replace(".", ",", 1), dict(ex)))
        if not numbering.parse_number(con, number, doc_type):
            out.append(Collision("number_format", f"Die Nummer {number} entspricht nicht dem Schema "
                                                  f"{'ANJJJJ-MMNN' if doc_type == 'angebot' else 'JJJJ-MMNN'}."))
    else:
        out.append(Collision("number_format", "Keine Belegnummer angegeben."))
    return out


def suggest_duplicate_number(con: sqlite3.Connection, number: str) -> str:
    i = 1
    while con.execute("SELECT 1 FROM document WHERE number=?", (f"{number}-D{i}",)).fetchone():
        i += 1
    return f"{number}-D{i}"


# ---------------------------------------------------------------- Übernahme
def import_parsed(con: sqlite3.Connection, inv: ParsedInvoice, *, number: str, doc_type: str, doc_date: str,
                  service_date: str | None, due_date: str | None, customer_id: int, lines: list[dict],
                  paid: bool = False, pay_date: str | None = None, tax_mode: str = "kleinunternehmer",
                  note: str = "", project_id: int | None = None) -> int:
    """Legt den Altbeleg festgeschrieben an, archiviert die Originaldatei (unverändert) und reserviert die Nummer."""
    if con.execute("SELECT 1 FROM document WHERE import_sha256=?", (inv.sha256,)).fetchone():
        raise ValueError("Diese Datei wurde bereits importiert.")
    if con.execute("SELECT 1 FROM document WHERE number=?", (number,)).fetchone():
        raise ValueError(f"Die Nummer {number} ist bereits vergeben.")
    year = (doc_date or "0000")[:4]
    target_dir = config.archive_dir() / "import" / year
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / inv.path.name
    k = 1
    while target.exists():
        target = target_dir / f"{inv.path.stem}_{k}{inv.path.suffix}"
        k += 1
    shutil.copy2(inv.path, target)
    try:
        target.chmod(0o444)
    except OSError:
        pass
    try:
        with transaction(con):
            doc = documents.new_draft(con, doc_type, customer_id, project_id)
            doc.update(doc_date=doc_date, service_date=service_date, due_date=due_date, tax_mode=tax_mode,
                       salutation_text=inv.salutation or doc["salutation_text"], zugferd=0)
            totals = documents.compute_totals([dict(l) for l in lines], tax_mode)
            doc_id = documents.save_draft(con, doc, lines)
            cust = repo.get_customer(con, customer_id)
            recipient = dict(cust)
            recipient["lines"] = inv.recipient_lines or repo.recipient_lines(cust)
            from app.db.database import company, dumps
            con.execute(
                """UPDATE document SET number=?, status='festgeschrieben', source='import', import_file_path=?,
                   import_sha256=?, import_note=?, recipient_json=?, company_json=?,
                   pdf_path=?, pdf_sha256=?, finalized_at=datetime('now','localtime') WHERE id=?""",
                (number, str(target), inv.sha256, note, dumps(recipient), dumps(company(con)),
                 str(target) if target.suffix.lower() == ".pdf" else None,
                 inv.sha256 if target.suffix.lower() == ".pdf" else None, doc_id))
            numbering.reserve(con, doc_type, number.split("-D")[0] if "-D" in number else number)
            audit.log(con, "import", "document", doc_id, {"number": number, "file": inv.path.name, "sha256": inv.sha256,
                                                          "gross": totals["gross"], "note": note})
            if paid and doc_type == "rechnung" and totals["gross"] > 0:
                con.execute("INSERT INTO payment(document_id, pay_date, amount_cents, method, note) VALUES (?,?,?,?,?)",
                            (doc_id, pay_date or doc_date, totals["gross"], "Überweisung", "Import: als bezahlt übernommen"))
                audit.log(con, "payment", "document", doc_id, {"date": pay_date or doc_date, "cents": totals["gross"],
                                                                "note": "Import"})
    except BaseException:
        try:
            target.chmod(0o666)
            target.unlink()
        except OSError:
            pass
        raise
    return doc_id
