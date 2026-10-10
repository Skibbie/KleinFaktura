"""Kundenimport aus CSV/XLSX mit frei wählbarer Spaltenzuordnung."""
from __future__ import annotations

import csv
import io
from pathlib import Path

TARGET_FIELDS = {
    "": "(ignorieren)",
    "salutation": "Anrede", "title": "Titel", "first_name": "Vorname", "last_name": "Nachname",
    "company": "Firma", "letter_salutation": "Briefanrede", "street": "Straße", "house_no": "Hausnummer",
    "zip": "PLZ", "city": "Ort", "country": "Land", "email": "E-Mail", "phone": "Telefon", "mobile": "Mobil",
    "vat_id": "USt-IdNr.", "notes": "Notiz", "customer_no": "Kundennummer",
}

_GUESS = {
    "anrede": "salutation", "titel": "title", "vorname": "first_name", "nachname": "last_name", "name": "last_name",
    "firma": "company", "firmenname": "company", "unternehmen": "company", "briefanrede": "letter_salutation",
    "straße": "street", "strasse": "street", "str": "street", "hausnummer": "house_no", "hnr": "house_no",
    "plz": "zip", "postleitzahl": "zip", "ort": "city", "stadt": "city", "land": "country", "e-mail": "email",
    "email": "email", "mail": "email", "telefon": "phone", "tel": "phone", "mobil": "mobile", "handy": "mobile",
    "ust-idnr": "vat_id", "notiz": "notes", "bemerkung": "notes", "kundennummer": "customer_no", "kdnr": "customer_no",
}


def read_table(path: Path | str) -> tuple[list[str], list[list[str]]]:
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = [["" if v is None else str(v).strip() for v in r] for r in ws.iter_rows(values_only=True)]
        wb.close()
    else:
        raw = path.read_bytes()
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
        except csv.Error:
            dialect = csv.excel
            dialect.delimiter = ";"
        rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), dialect)]
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        return [], []
    return rows[0], rows[1:]


def guess_mapping(headers: list[str]) -> list[str]:
    out = []
    for h in headers:
        key = h.strip().lower().rstrip(".:").replace("_", "-")
        out.append(_GUESS.get(key, _GUESS.get(key.replace("-", ""), "")))
    return out


def rows_to_customers(rows: list[list[str]], mapping: list[str]) -> list[dict]:
    result = []
    for r in rows:
        c: dict = {}
        for i, field in enumerate(mapping):
            if field and i < len(r) and r[i]:
                c[field] = (c.get(field, "") + " " + r[i]).strip() if field in c else r[i]
        if c.get("house_no"):
            c["street"] = f"{c.get('street', '')} {c.pop('house_no')}".strip()
        c.pop("house_no", None)
        if c.get("last_name") or c.get("company"):
            c["source"] = "csv"
            result.append(c)
    return result
