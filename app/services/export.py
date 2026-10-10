"""Exporte: Umsatzliste, Einnahmen (EÜR), DATEV-Buchungsstapel, GoBD-Datenexport für die Betriebsprüfung."""
from __future__ import annotations

import csv
import io
import sqlite3
import zipfile
from datetime import datetime
from pathlib import Path

from app import config
from app.db import repo
from app.db.database import get_setting
from app.money import cents_to_decimal, fmt_date
from app.services import documents

DOC_COLS = ["Belegnummer", "Belegart", "Datum", "Leistungsdatum", "Fällig", "Kundennr.", "Kunde", "Netto", "USt",
            "Brutto", "Bezahlt", "Offen", "Status", "Quelle"]


def _customer_name(con, cid):
    return repo.customer_display_name(repo.get_customer(con, cid) or {}) if cid else ""


def revenue_rows(con: sqlite3.Connection, year: int | None) -> list[list]:
    rows = []
    for d in documents.list_documents(con, ("rechnung", "storno"), year=year):
        if d["status"] != "festgeschrieben":
            continue
        rows.append([d["number"], documents.DOC_TYPE_LABEL[d["doc_type"]], fmt_date(d["doc_date"]),
                     fmt_date(d["service_date"]), fmt_date(d["due_date"]), d.get("customer_no") or "",
                     _customer_name(con, d["customer_id"]), cents_to_decimal(d["total_net_cents"]),
                     cents_to_decimal(d["total_tax_cents"]), cents_to_decimal(d["total_gross_cents"]),
                     cents_to_decimal(d["paid_cents"]), cents_to_decimal(d["open_cents"]), d["status_label"],
                     "Import" if d["source"] == "import" else config.APP_NAME])
    rows.sort(key=lambda r: r[0])
    return rows


def income_rows(con: sqlite3.Connection, year: int | None) -> list[list]:
    sql = """SELECT p.pay_date, p.amount_cents, p.method, p.note, d.number, d.customer_id
             FROM payment p JOIN document d ON d.id=p.document_id"""
    args = []
    if year:
        sql += " WHERE substr(p.pay_date,1,4)=?"
        args.append(f"{year:04d}")
    sql += " ORDER BY p.pay_date, p.id"
    return [[fmt_date(r["pay_date"]), r["number"], _customer_name(con, r["customer_id"]),
             cents_to_decimal(r["amount_cents"]), r["method"], r["note"]] for r in con.execute(sql, args)]


def write_xlsx(path: Path, sheets: dict[str, tuple[list[str], list[list]]]) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)
    for name, (header, rows) in sheets.items():
        ws = wb.create_sheet(name[:31])
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([float(v) if hasattr(v, "quantize") else v for v in r])
        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in col) + 2
            ws.column_dimensions[col[0].column_letter].width = min(width, 50)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, float):
                    c.number_format = '#,##0.00 "€"'
    wb.save(path)
    return path


def write_csv(path: Path, header: list[str], rows: list[list]) -> Path:
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(header)
        for r in rows:
            w.writerow([f"{v:.2f}".replace(".", ",") if hasattr(v, "quantize") else v for v in r])
    return path


def export_revenue(con: sqlite3.Connection, path: Path, year: int | None) -> Path:
    path = Path(path)
    if path.suffix.lower() == ".csv":
        return write_csv(path, DOC_COLS, revenue_rows(con, year))
    return write_xlsx(path, {
        "Umsatzliste": (DOC_COLS, revenue_rows(con, year)),
        "Einnahmen (Zahlungen)": (["Datum", "Rechnung", "Kunde", "Betrag", "Art", "Notiz"], income_rows(con, year)),
    })


def export_datev(con: sqlite3.Connection, path: Path, year: int) -> Path:
    """DATEV-Buchungsstapel (EXTF, Version 700) – Ist-Prinzip: Zahlungseingänge Bank an Erlöse."""
    bank = get_setting(con, "datev.konto_bank", "1200")
    erloes_ku = get_setting(con, "datev.konto_erloes", "8195")
    berater = get_setting(con, "datev.berater", "1001")
    mandant = get_setting(con, "datev.mandant", "1")
    now = datetime.now()
    header = ["EXTF", "700", "21", "Buchungsstapel", "13", now.strftime("%Y%m%d%H%M%S000"), "", "EF", "",
              "", berater, mandant, f"{year}0101", "4", f"{year}0101", f"{year}1231", config.APP_NAME, "", "1", "0",
              "0", "EUR"]
    cols = ["Umsatz (ohne Soll/Haben-Kz)", "Soll/Haben-Kennzeichen", "WKZ Umsatz", "Kurs", "Basis-Umsatz",
            "WKZ Basis-Umsatz", "Konto", "Gegenkonto (ohne BU-Schlüssel)", "BU-Schlüssel", "Belegdatum",
            "Belegfeld 1", "Belegfeld 2", "Skonto", "Buchungstext"]
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", quoting=csv.QUOTE_MINIMAL)
    w.writerow(header)
    w.writerow(cols)
    sql = """SELECT p.pay_date, p.amount_cents, d.number, d.customer_id FROM payment p
             JOIN document d ON d.id=p.document_id WHERE substr(p.pay_date,1,4)=? ORDER BY p.pay_date, p.id"""
    for r in con.execute(sql, (f"{year:04d}",)):
        amt = r["amount_cents"]
        name = _customer_name(con, r["customer_id"])
        w.writerow([f"{abs(amt) / 100:.2f}".replace(".", ","), "S" if amt > 0 else "H", "EUR", "", "", "",
                    bank, erloes_ku, "", datetime.strptime(r["pay_date"], "%Y-%m-%d").strftime("%d%m"),
                    r["number"][:36], "", "", f"Zahlung {r['number']} {name}"[:60]])
    Path(path).write_bytes(buf.getvalue().encode("cp1252", errors="replace"))
    return Path(path)


def export_gobd(con: sqlite3.Connection, path: Path) -> Path:
    """Datenexport für die Betriebsprüfung (Datenträgerüberlassung): alle Tabellen als CSV + Beschreibung."""
    tables = {
        "belege.csv": "SELECT id, doc_type, number, status, customer_id, project_id, doc_date, service_date, "
                      "service_date_to, due_date, tax_mode, total_net_cents, total_tax_cents, total_gross_cents, "
                      "ref_document_id, source, pdf_sha256, xml_sha256, import_sha256, finalized_at FROM document "
                      "WHERE status='festgeschrieben' ORDER BY number",
        "positionen.csv": "SELECT l.document_id, d.number, l.pos, l.title, l.description, l.quantity, l.unit, "
                          "l.unit_price_cents, l.tax_rate, l.total_cents FROM document_line l JOIN document d "
                          "ON d.id=l.document_id WHERE d.status='festgeschrieben' ORDER BY d.number, l.pos",
        "zahlungen.csv": "SELECT p.id, d.number, p.pay_date, p.amount_cents, p.method, p.note, p.created_at "
                         "FROM payment p JOIN document d ON d.id=p.document_id ORDER BY p.id",
        "mahnungen.csv": "SELECT m.id, d.number, m.level, m.dunning_date, m.new_due_date, m.fee_cents, m.open_cents, "
                         "m.pdf_sha256 FROM dunning m JOIN document d ON d.id=m.document_id ORDER BY m.id",
        "kunden.csv": "SELECT id, customer_no, salutation, title, first_name, last_name, company, street, zip, city, "
                      "country, vat_id FROM customer ORDER BY id",
        "ausgaben.csv": "SELECT id, exp_date, supplier, description, category, amount_cents, vat_rate, km, useful_life, "
                        "payment_method, receipt_sha256, source, created_at, updated_at FROM expense ORDER BY exp_date, id",
        "fahrten.csv": "SELECT id, trip_date, customer_id, project_id, document_id, destination, purpose, km_one_way, "
                       "round_trip, km, rate, amount_cents, note, created_at, updated_at FROM trip ORDER BY trip_date, id",
        "protokoll.csv": "SELECT id, ts, action, entity, entity_id, data, prev_hash, hash FROM audit_log ORDER BY id",
    }
    readme = [f"GoBD-Datenexport {config.APP_NAME}", f"Erstellt: {datetime.now():%d.%m.%Y %H:%M}", "",
              "Zeichensatz UTF-8, Trennzeichen Semikolon, Beträge in Cent (Ganzzahl).", ""]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, sql in tables.items():
            cur = con.execute(sql)
            buf = io.StringIO()
            w = csv.writer(buf, delimiter=";")
            cols = [c[0] for c in cur.description]
            w.writerow(cols)
            n = 0
            for r in cur:
                w.writerow(list(r))
                n += 1
            z.writestr(name, buf.getvalue().encode("utf-8-sig"))
            readme.append(f"{name}: {n} Datensätze; Spalten: {', '.join(cols)}")
        z.writestr("LIESMICH.txt", "\n".join(readme))
    return Path(path)
